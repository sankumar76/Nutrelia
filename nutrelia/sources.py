"""Load public food databases into Nutrelia's public list.

- INDB, the Indian Nutrient Databank (https://www.anuvaad.org.in/indian-nutrient-databank/):
  the .xlsx (or a .csv saved from it). Values are per 100 g.
- USDA FoodData Central (https://fdc.nal.usda.gov/download-datasets): the "SR Legacy" and
  "Foundation Foods" CSV downloads (.zip). Values are per 100 g.
- Open Food Facts (https://world.openfoodfacts.org): packaged products, either fetched from
  its search API (for example every product sold in India) or read from its CSV export.

Every loader yields PublicFood rows; `save` writes them to the database. Loading the same
source again updates the foods already there instead of adding copies.
"""
from __future__ import annotations

import csv
import gzip
import io
import json
import re
import sqlite3
import time
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator

from .db import save_public_food
from .units import canonical_unit, count_unit

USER_AGENT = "Nutrelia/1.0 (https://github.com/sankumar76/Nutrelia)"


@dataclass
class PublicFood:
    source: str  # indb, usda, off
    ref: str  # the food's id in that database
    name: str
    nutrients: dict = field(default_factory=dict)
    base_amount: float = 100
    unit: str = "g"
    portions: list = field(default_factory=list)  # (unit, grams, label): household measures


MAX_PORTION_G = 600
_MASS_WORDS = {"g", "kg", "oz", "lb", "mg"}


def portion_unit(label: str) -> str | None:
    """The one word people would type for a household measure: "soup bowl" -> bowl,
    "cup, chopped" -> cup, "medium (7-1/4" long)" -> medium, "tablespoon" -> tbsp."""
    head = re.split(r"[,(]", label or "", maxsplit=1)[0]
    words = re.findall(r"[a-zA-Z]+", head.lower())
    if not words:
        return None
    word = words[-1]
    unit = canonical_unit(word) or count_unit(word)
    return None if unit in _MASS_WORDS else unit


def save(conn: sqlite3.Connection, foods: Iterable[PublicFood]) -> dict:
    added = updated = skipped = 0
    with conn:
        for f in foods:
            if not f.name:
                skipped += 1
                continue
            n = {k: (None if v is None else round(v, 2)) for k, v in f.nutrients.items()}
            if n.get("calories") is None and None not in (n.get("carbs"), n.get("protein"), n.get("fat")):
                n["calories"] = round(4 * n["carbs"] + 4 * n["protein"] + 9 * n["fat"], 1)
            if n.get("calories") is None:
                skipped += 1
                continue
            # a "dish", "jar" or whole chicken is a recipe's yield, not something one person eats
            portions = [(u, round(g, 1), label) for u, g, label in f.portions if u and 0 < g <= MAX_PORTION_G]
            if save_public_food(conn, f.source, f.ref, f.name[:200], f.base_amount, f.unit, n, portions):
                added += 1
            else:
                updated += 1
    return {"added": added, "updated": updated, "skipped": skipped}


def _num(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    v = str(v).strip().replace(",", "")
    try:
        return float(v)
    except ValueError:
        return None


# --- INDB ------------------------------------------------------------------

# INDB column -> Nutrelia nutrient; the first column found wins
_INDB_COLUMNS = {
    "calories": ["energy_kcal", "energy (kcal)", "energy"],
    "carbs": ["carb_g", "carbohydrate_g", "carbohydrates_g", "carbohydrate"],
    "fat": ["fat_g", "total_fat_g", "fat"],
    "protein": ["protein_g", "protein"],
    "fiber": ["fibre_g", "fiber_g", "dietary_fibre_g", "fibre", "fiber"],
    "sugar": ["freesugar_g", "sugar_g", "sugars_g", "total_sugar_g"],
    "sodium": ["sodium_mg", "sodium"],
    "cholesterol": ["cholesterol_mg", "cholesterol"],
}


def _rows_from_table(path: Path) -> Iterator[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:  # read-only workbooks keep the file open until closed
            header = None
            for row in wb.worksheets[0].iter_rows(values_only=True):
                cells = [str(c).strip().lower() if c is not None else "" for c in row]
                if header is None:
                    if any("name" in c for c in cells):
                        header = cells
                    continue
                yield dict(zip(header, row))
        finally:
            wb.close()
    else:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            for row in csv.DictReader(fh):
                yield {(k or "").strip().lower(): v for k, v in row.items()}


def load_indb(path: str | Path) -> Iterator[PublicFood]:
    path = Path(path)
    cols = None
    for i, row in enumerate(_rows_from_table(path)):
        if cols is None:
            keys = list(row)
            name_col = next((k for k in ("food_name", "name", "food name", "recipe_name") if k in keys),
                            next((k for k in keys if "name" in k), None))
            ref_col = next((k for k in ("food_code", "code", "recipe_code") if k in keys), None)
            serving_col = "servings_unit" if "servings_unit" in keys else None
            cols = {"name": name_col, "ref": ref_col}
            for n, options in _INDB_COLUMNS.items():
                cols[n] = next((k for k in options if k in keys), None)
            if not cols["name"] or not cols["calories"]:
                raise ValueError("This doesn't look like the INDB table: no food name or energy (kcal) column found.")
        name = str(row.get(cols["name"]) or "").strip()
        if not name:
            continue
        ref = str(row.get(cols["ref"]) or "").strip() if cols["ref"] else ""
        food = PublicFood("indb", ref or name.lower(), name,
                          {n: _num(row.get(c)) for n, c in cols.items() if n not in ("name", "ref") and c})
        if serving_col and row.get(serving_col):
            # INDB gives one household serving ("bowl") with its own nutrition; its weight is
            # the ratio of the serving's values to the per-100 g ones
            label = str(row[serving_col]).strip()
            for key in ("energy_kcal", "carb_g", "protein_g", "fat_g"):
                per100, per_serving = _num(row.get(key)), _num(row.get(f"unit_serving_{key}"))
                if per100 and per_serving:
                    food.portions.append((portion_unit(label), 100 * per_serving / per100, label))
                    break
        yield food


# --- USDA FoodData Central ---------------------------------------------------

# nutrient numbers (nutrient.csv "nutrient_nbr"); energy prefers 208, then the Atwater values
_USDA_NUMBERS = {"203": "protein", "204": "fat", "205": "carbs", "291": "fiber", "269": "sugar",
                 "307": "sodium", "601": "cholesterol", "208": "calories"}
_USDA_ENERGY_FALLBACK = {"957": 1, "958": 2}  # Atwater general / specific, used when 208 is missing
_USDA_TYPES = {"sr_legacy_food", "foundation_food", "survey_fndds_food"}


def _open_member(z: zipfile.ZipFile, name: str):
    member = next((m for m in z.namelist() if m.rsplit("/", 1)[-1] == name), None)
    if member is None:
        raise ValueError(f"{name} is missing from the USDA download.")
    return io.TextIOWrapper(z.open(member), encoding="utf-8", newline="")


def load_usda(path: str | Path) -> Iterator[PublicFood]:
    """Foods from a FoodData Central CSV download (SR Legacy, Foundation Foods or FNDDS)."""
    with zipfile.ZipFile(path) as z:
        wanted = {}  # nutrient id -> our nutrient (or an energy fallback key)
        for r in csv.DictReader(_open_member(z, "nutrient.csv")):
            nbr = (r.get("nutrient_nbr") or "").split(".")[0].strip()
            if nbr in _USDA_NUMBERS:
                wanted[r["id"]] = _USDA_NUMBERS[nbr]
            elif nbr in _USDA_ENERGY_FALLBACK:
                wanted[r["id"]] = f"energy{_USDA_ENERGY_FALLBACK[nbr]}"
        foods = {}
        for r in csv.DictReader(_open_member(z, "food.csv")):
            if r.get("data_type") in _USDA_TYPES:
                foods[r["fdc_id"]] = {"name": r["description"].strip()}
        for r in csv.DictReader(_open_member(z, "food_nutrient.csv")):
            food = foods.get(r["fdc_id"])
            key = wanted.get(r["nutrient_id"])
            if food is not None and key:
                food[key] = _num(r.get("amount"))
        try:
            measures = {r["id"]: r["name"] for r in csv.DictReader(_open_member(z, "measure_unit.csv"))}
            portion_rows = list(csv.DictReader(_open_member(z, "food_portion.csv")))
        except ValueError:
            measures, portion_rows = {}, []
        portion_rows.sort(key=lambda r: (r["fdc_id"], _num(r.get("seq_num")) or 0, _num(r.get("id")) or 0))
        for r in portion_rows:
            food = foods.get(r["fdc_id"])
            grams, amount = _num(r.get("gram_weight")), _num(r.get("amount")) or 1
            if food is None or not grams:
                continue
            measure = measures.get(r.get("measure_unit_id"), "")
            label = (r.get("modifier") or "").strip() if measure in ("", "undetermined") else \
                " ".join(x for x in (measure, (r.get("modifier") or "").strip()) if x)
            if r.get("portion_description"):
                label = label or r["portion_description"].strip()
            food.setdefault("portions", []).append((portion_unit(label), grams / amount, label))
    for fdc_id, f in foods.items():
        if f.get("calories") is None:
            f["calories"] = f.get("energy1") if f.get("energy1") is not None else f.get("energy2")
        yield PublicFood("usda", fdc_id, f["name"],
                         {n: f.get(n) for n in ("calories", "carbs", "fat", "protein", "fiber", "sugar",
                                                "sodium", "cholesterol")}, portions=f.get("portions", []))


# --- Open Food Facts ---------------------------------------------------------

OFF_FIELDS = "code,product_name,product_name_en,brands,nutriments,countries_tags,serving_size,serving_quantity"


def _off_food(p: dict) -> PublicFood | None:
    n = p.get("nutriments") or {}

    def g(key):
        return _num(n.get(f"{key}_100g", p.get(f"{key}_100g")))

    name = (p.get("product_name_en") or p.get("product_name") or "").strip()
    code = str(p.get("code") or "").strip()
    if not name or not code:
        return None
    brand = (p.get("brands") or "").split(",")[0].strip()
    if brand and brand.lower() not in name.lower():
        name = f"{brand} - {name}"
    kcal = g("energy-kcal")
    if kcal is None and g("energy") is not None:
        kcal = g("energy") / 4.184  # kJ
    sodium, chol = g("sodium"), g("cholesterol")
    serving = _num(p.get("serving_quantity"))
    portions = [("serving", serving, p.get("serving_size") or "serving")] if serving else []
    return PublicFood("off", code, re.sub(r"\s+", " ", name), portions=portions, nutrients={
        "calories": None if kcal is None else round(kcal, 1),
        "carbs": g("carbohydrates"), "fat": g("fat"), "protein": g("proteins"),
        "fiber": g("fiber"), "sugar": g("sugars"),
        "sodium": None if sodium is None else round(sodium * 1000, 1),  # g -> mg
        "cholesterol": None if chol is None else round(chol * 1000, 1),
    })


def _get_json(url: str) -> dict:
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.load(resp)


def search_off(query: str, country: str | None = None, page_size: int = 24) -> list[PublicFood]:
    """One search of Open Food Facts by name (used by the admin page)."""
    params = {"search_terms": query, "search_simple": 1, "action": "process", "json": 1,
              "page_size": page_size, "fields": OFF_FIELDS}
    if country:
        params.update({"tagtype_0": "countries", "tag_contains_0": "contains", "tag_0": country})
    data = _get_json("https://world.openfoodfacts.org/cgi/search.pl?" + urllib.parse.urlencode(params))
    return [f for f in map(_off_food, data.get("products", [])) if f]


def fetch_off_country(country: str = "india", limit: int | None = None, page_size: int = 100,
                      pause: float = 6.5, log=print) -> Iterator[PublicFood]:
    """Every product Open Food Facts lists for a country. Search allows about 10 requests a minute."""
    page, seen = 1, 0
    while True:
        params = {"countries_tags_en": country, "fields": OFF_FIELDS, "page_size": page_size, "page": page,
                  "sort_by": "unique_scans_n"}
        data = _get_json("https://world.openfoodfacts.org/api/v2/search?" + urllib.parse.urlencode(params))
        products = data.get("products", [])
        if page == 1:
            log(f"Open Food Facts lists {data.get('count', '?')} products for {country}.")
        for p in products:
            f = _off_food(p)
            if f:
                yield f
                seen += 1
                if limit and seen >= limit:
                    return
        if len(products) < page_size:
            return
        page += 1
        log(f"  page {page}...")
        time.sleep(pause)


def load_off_csv(path: str | Path, country: str | None = "india") -> Iterator[PublicFood]:
    """Products from Open Food Facts' CSV export (tab separated, optionally .gz)."""
    path = Path(path)
    csv.field_size_limit(1 << 30)
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8", newline="") as fh:
        for row in csv.DictReader(fh, delimiter="\t"):
            if country and f"en:{country}" not in (row.get("countries_tags") or ""):
                continue
            f = _off_food(row)
            if f:
                yield f
